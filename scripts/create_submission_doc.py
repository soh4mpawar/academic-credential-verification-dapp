import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Page Margins - 0.65 in
for section in doc.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink {nsdecls("w")} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}">'
        f'<w:r><w:rPr><w:color w:val="2563EB"/><w:u w:val="single"/><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/></w:rPr>'
        f'<w:t>{text}</w:t></w:r></w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

# ----------------------------------------------------------------------
# PAGE 1: Header, Metadata, Problem Statement, Architecture
# ----------------------------------------------------------------------
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_after = Pt(2)
run_inst = p_title.add_run("MANIPAL INSTITUTE OF TECHNOLOGY\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
run_inst.font.name = "Calibri"
run_inst.font.size = Pt(12)
run_inst.font.bold = True
run_inst.font.color.rgb = RGBColor(15, 23, 42)

run_sub = p_title.add_run("BLOCKCHAIN TECHNOLOGY (CSE4444) — ASSIGNMENT 2\n")
run_sub.font.name = "Calibri"
run_sub.font.size = Pt(14)
run_sub.font.bold = True
run_sub.font.color.rgb = RGBColor(37, 99, 235)

p_rule = doc.add_paragraph()
p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_rule.paragraph_format.space_after = Pt(6)
r_rule = p_rule.add_run("―" * 62)
r_rule.font.color.rgb = RGBColor(203, 213, 225)

# Student Metadata Table
table = doc.add_table(rows=5, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False

metadata = [
    ("Student Name:", "Soham Sandeep Pawar", False),
    ("Registration Number:", "230905312", False),
    ("Assignment Title:", "Design and Development of an Original DApp (Assignment 2)", False),
    ("Title of Developed DApp:", "Decentralized Academic Credential Verification System (EduVerify)", False),
    ("GitHub Repository Link:", "https://github.com/soh4mpawar/academic-credential-verification-dapp", True)
]

for idx, (label, val, is_link) in enumerate(metadata):
    row = table.rows[idx]
    
    # Col 0: Label
    cell0 = row.cells[0]
    cell0.width = Inches(2.0)
    p0 = cell0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r0 = p0.add_run(label)
    r0.font.name = "Calibri"
    r0.font.size = Pt(9.5)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(30, 41, 59)
    set_cell_background(cell0, "F1F5F9")
    set_cell_margins(cell0, 50, 50, 80, 80)
    
    # Col 1: Value
    cell1 = row.cells[1]
    cell1.width = Inches(5.2)
    p1 = cell1.paragraphs[0]
    if is_link:
        add_hyperlink(p1, val, val)
    else:
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        if "Soham" in val or "230905312" in val or "EduVerify" in val:
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r1.font.color.rgb = RGBColor(51, 65, 85)
    set_cell_background(cell1, "FFFFFF")
    set_cell_margins(cell1, 50, 50, 80, 80)

# Section 1
h1 = doc.add_heading("1. Problem Description & Real-World Motivation", level=2)
h1.style.font.color.rgb = RGBColor(30, 41, 59)
h1.paragraph_format.space_before = Pt(8)
h1.paragraph_format.space_after = Pt(2)
h1.paragraph_format.keep_with_next = True

p_desc = doc.add_paragraph(
    "Academic credential fraud, including counterfeit degree certificates and unverified grade submissions, "
    "presents a severe problem for universities, employers, and professional licensing authorities. Traditional "
    "verification workflows require manual inquiries, letters of confirmation, or centralized third-party registries, "
    "all of which suffer from significant delays, administrative expenses, and single-point-of-failure vulnerabilities.\n"
    "EduVerify resolves this by implementing a decentralized academic credential registry on the Ethereum blockchain. "
    "Accredited universities and departments mint immutable academic certificates directly onto the ledger. "
    "Each certificate is identified by a unique cryptographic identifier and records the student's name, roll number, "
    "degree title, cumulative grade, issuing authority, and timestamp. Verifiers and prospective employers can instantly "
    "validate credential authenticity directly against smart contract state without intermediary fees or delays."
)
p_desc.style.font.name = "Calibri"
p_desc.style.font.size = Pt(9.5)
p_desc.paragraph_format.space_after = Pt(6)

# Section 2
h2 = doc.add_heading("2. Smart Contract & Frontend Architecture", level=2)
h2.style.font.color.rgb = RGBColor(30, 41, 59)
h2.paragraph_format.space_before = Pt(6)
h2.paragraph_format.space_after = Pt(2)
h2.paragraph_format.keep_with_next = True

p_arch = doc.add_paragraph(
    "• Smart Contract (AcademicCredential.sol): Implemented in Solidity ^0.8.20 and compiled targeting the Shanghai EVM. "
    "Maintains access-controlled registries for accredited issuers and graduate records. Incorporates input validations "
    "rejecting empty parameters, duplicate certificate IDs, unauthorized minting, and unauthorized revocations.\n"
    "• Frontend Interface: Developed using HTML5, CSS3, and JavaScript, providing dedicated workflows for Public "
    "Certificate Verification, Department Credential Issuance, Revocation Management, and Administrative Issuer Control.\n"
    "• Web3 & MetaMask Integration: Implemented using Ethers.js connecting via window.ethereum to the local Ganache "
    "blockchain network (Chain ID: 1337) for transaction signing, account detection, and real-time event synchronization."
)
p_arch.style.font.name = "Calibri"
p_arch.style.font.size = Pt(9.5)
p_arch.paragraph_format.space_after = Pt(6)

# ----------------------------------------------------------------------
# PAGE 2: Test Cases Table (Explicit Page Break, No Split)
# ----------------------------------------------------------------------
doc.add_page_break()

h3 = doc.add_heading("3. Test Cases & Live EVM Verification Results", level=2)
h3.style.font.color.rgb = RGBColor(30, 41, 59)
h3.paragraph_format.space_before = Pt(4)
h3.paragraph_format.space_after = Pt(4)
h3.paragraph_format.keep_with_next = True

p_test_desc = doc.add_paragraph(
    "The Solidity smart contract was compiled using solc v0.8.20 (Shanghai) and evaluated through automated EVM execution "
    "against a live Ganache blockchain node. All 12 test assertions covering Valid, Invalid, and Unauthorised inputs "
    "executed directly against contract bytecode with verified transaction receipts and EVM revert exceptions:"
)
p_test_desc.style.font.name = "Calibri"
p_test_desc.style.font.size = Pt(9.5)
p_test_desc.paragraph_format.space_after = Pt(6)
p_test_desc.paragraph_format.keep_with_next = True

# 12 Test Cases matching real EVM test output
test_rows_data = [
    ("TC-01", "Valid", "Constructor sets deployer as owner and initial admin", "Owner address recorded, IssuerAdded event emitted", "PASS"),
    ("TC-02", "Valid", "Owner authorizes CSE Department (0x6174...)", "Tx: 0x5ced633e... | Gas: 139,874 gas", "PASS"),
    ("TC-03", "Valid", "Authorized issuer mints cert (CERT-2026-CSE-001)", "Tx: 0xdc044036... | Gas: 280,516 gas (Block #3)", "PASS"),
    ("TC-04", "Valid", "Public query of valid Certificate ID from EVM storage", "Returns verified record: Soham Sandeep Pawar (230905312)", "PASS"),
    ("TC-05", "Valid", "Department revokes certificate with valid reason", "Tx: 0x2218a3da... | isRevoked set to true", "PASS"),
    ("TC-06", "Invalid", "Issue certificate with empty student name string (\"\")", "EVM Revert: 'Student name cannot be empty'", "PASS"),
    ("TC-07", "Invalid", "Issue certificate with duplicate Certificate ID", "EVM Revert: 'Certificate with this ID already exists'", "PASS"),
    ("TC-08", "Invalid", "Query non-existent Certificate ID", "Returns exists: false gracefully without EVM failure", "PASS"),
    ("TC-09", "Invalid", "Revoke certificate with empty reason string (\"\")", "EVM Revert: 'Revocation reason cannot be empty'", "PASS"),
    ("TC-10", "Unauth", "Non-owner attempts to call addIssuer(...)", "EVM Revert: 'Caller is not the contract owner'", "PASS"),
    ("TC-11", "Unauth", "Non-authorized address attempts to issue certificate", "EVM Revert: 'Caller is not authorized to issue certificates'", "PASS"),
    ("TC-12", "Unauth", "Unauthorized account attempts to revoke certificate", "EVM Revert: 'Only issuing authority or owner can revoke'", "PASS")
]

test_table = doc.add_table(rows=len(test_rows_data) + 1, cols=5)
test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
test_table.autofit = False

# Table header formatting
headers = ["Test ID", "Category", "Description / Input", "Observed EVM Behavior / Receipt", "Status"]
header_row = test_table.rows[0]
trPr = header_row._tr.get_or_add_trPr()
trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>')) # Repeat header if ever needed
trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

col_widths = [Inches(0.75), Inches(0.8), Inches(2.6), Inches(2.45), Inches(0.6)]
for idx, text in enumerate(headers):
    cell = header_row.cells[idx]
    cell.width = col_widths[idx]
    set_cell_background(cell, "1E293B")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_margins(cell, 40, 40, 50, 50)

# Table rows
for row_idx, data in enumerate(test_rows_data):
    row = test_table.rows[row_idx + 1]
    row_trPr = row._tr.get_or_add_trPr()
    row_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>')) # Prevent row split
    
    bg_color = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
    for col_idx, text in enumerate(data):
        cell = row.cells[col_idx]
        cell.width = col_widths[col_idx]
        set_cell_background(cell, bg_color)
        p = cell.paragraphs[0]
        if col_idx in [0, 1, 4]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(8)
        if text == "PASS":
            r.font.bold = True
            r.font.color.rgb = RGBColor(22, 163, 74)
        elif "Revert" in text:
            r.font.color.rgb = RGBColor(220, 38, 38)
        elif "Tx:" in text:
            r.font.color.rgb = RGBColor(37, 99, 235)
        else:
            r.font.color.rgb = RGBColor(30, 41, 59)
        set_cell_margins(cell, 30, 30, 40, 40)

# Summary paragraph under table
p_summary = doc.add_paragraph()
p_summary.paragraph_format.space_before = Pt(6)
r_sum = p_summary.add_run("Automated EVM Test Suite Result: 12 / 12 Tests Passed (100% Success Rate) against live contract bytecode.")
r_sum.font.name = "Calibri"
r_sum.font.size = Pt(9)
r_sum.font.bold = True
r_sum.font.color.rgb = RGBColor(22, 163, 74)

# ----------------------------------------------------------------------
# Helper for Adding Figure with Non-Breaking Caption
# ----------------------------------------------------------------------
def add_figure(title, img_path, is_page_start=False):
    if is_page_start:
        doc.add_page_break()
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.keep_with_next = True # CRITICAL: Caption NEVER separates from image!
    
    r_title = p_title.add_run(title)
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 41, 59)
    
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(6)
        p_img.paragraph_format.keep_with_next = True
        p_img.add_run().add_picture(img_path, width=Inches(5.8))

# ----------------------------------------------------------------------
# PAGES 3, 4, 5: Visual Evidence Figures (2 per page, zero orphan captions)
# ----------------------------------------------------------------------
doc.add_page_break()
h4 = doc.add_heading("4. Visual Evidence & Blockchain Execution Captures", level=2)
h4.style.font.color.rgb = RGBColor(30, 41, 59)
h4.paragraph_format.space_before = Pt(2)
h4.paragraph_format.space_after = Pt(2)
h4.paragraph_format.keep_with_next = True

# Page 3: Figures 1 & 2
add_figure(
    "Figure 1: Smart Contract Compilation in Remix IDE (Solidity v0.8.20)",
    "screenshots/01_smart_contract_compilation.png"
)
add_figure(
    "Figure 2: Smart Contract Deployment in Remix IDE Environment (Block #1)",
    "screenshots/02_smart_contract_deployment.png"
)

# Page 4: Figures 3 & 4
add_figure(
    "Figure 3: MetaMask Wallet Connection & Active Role Detection in EduVerify UI",
    "screenshots/03_metamask_wallet_connection.png",
    is_page_start=True
)
add_figure(
    "Figure 4: Successful Credential Issuance Transaction Execution on Blockchain (Block #3)",
    "screenshots/04_successful_transaction.png"
)

# Page 5: Figures 5 & 6 (Figure 5 caption is guaranteed to stay with Figure 5 image!)
add_figure(
    "Figure 5: Failed Transaction Handling (Solidity require Revert & EVM Exception Display)",
    "screenshots/05_failed_transaction.png",
    is_page_start=True
)
add_figure(
    "Figure 6: Verified Blockchain Credential Retrieval & Display Interface",
    "screenshots/06_blockchain_data_displayed_ui.png"
)

output_path = "SohamPawar_230905312_BT_Assignment2.docx"
doc.save(output_path)
print(f"Document successfully created: {output_path}")
